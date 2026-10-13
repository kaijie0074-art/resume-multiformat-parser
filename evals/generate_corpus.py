"""合成语料生成器：35 份简历 + 字段级 ground truth。

全部程序合成、固定随机种子（20261013），无任何真实个人信息。

构成：
- 10 份单栏文本 PDF（single-01..10）
- 10 份双栏文本 PDF（two-col-01..10）：左栏时间轴+教育+证书，右栏经历——
  线性抽取必然左右交错（基线缺陷 RPI-001 的复现源）
- 10 份图片型 PDF（scanned-01..10）：PIL 渲染文本为位图再嵌页，无文本层
- 5 份 docx 表格简历（docx-01..05）：段落 + 教育表格（含合并单元格）

输出：
- evals/corpus/<stem>.{pdf,docx}
- evals/ground_truth/<stem>.json

用法：python3 evals/generate_corpus.py [--outdir evals]
"""

from __future__ import annotations

import json
import random
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

# ---------------- 合成事实池 ----------------

NAMES = ["王思远", "林晓阳", "陈雨桐", "赵启明", "周子墨", "吴悦宁", "郑天翊", "孙艺涵", "钱嘉树", "冯静姝"]
PHONES = [f"138{i:08d}" for i in range(10)]
EMAILS = [f"user{i}@example.com" for i in range(10)]
SCHOOLS = ["沪江大学", "东南文理学院", "临港职业技术学院", "青溪大学", "白鹭学院"]
DEGREES = ["本科", "本科", "大专"]
COMPANIES = [
    "星桥信息技术有限公司", "澄海贸易有限公司", "木棉传媒有限公司", "岚山咨询有限公司", "汀洲网络科技有限公司",
    "云汀数据服务有限公司", "栗山互联网有限公司", "南浦电子商务有限公司", "鹿鸣科技有限公司", "苇岸智能科技有限公司",
]
TITLES = ["运营实习生", "后端开发实习生", "数据分析实习生", "产品助理", "测试实习生"]
DESCRIPTIONS = [
    "参与订单系统开发，日均处理 3000 单",
    "维护 200 个用户社群，月均转化 150 单",
    "搭建业务报表，覆盖 120 项核心指标",
    "参与小程序迭代，服务 800 名种子用户",
    "协助策划 6 场校园活动，触达 5000 人次",
]
CERTS_POOL = ["计算机二级", "普通话二甲", "驾驶证", "CET-4"]
EDU_DATES = ["2022.09-2026.06", "2023.09-2026.06", "2021.09-2024.06"]
WORK_DATES = ["2024.07-2025.01", "2023.09-2024.06", "2025.03-2025.08"]


def make_profile(idx: int, rng: random.Random) -> dict:
    certs = rng.sample(CERTS_POOL, k=rng.choice([1, 2]))
    n_work = rng.choice([1, 2])
    works = []
    for w in range(n_work):
        works.append({
            "company": COMPANIES[(idx * 2 + w) % len(COMPANIES)],
            "title": TITLES[(idx + w) % len(TITLES)],
            "date": WORK_DATES[(idx + w) % len(WORK_DATES)],
            "desc": DESCRIPTIONS[(idx * 3 + w) % len(DESCRIPTIONS)],
        })
    return {
        "name": NAMES[idx % len(NAMES)],
        "phone": PHONES[idx % len(PHONES)],
        "email": EMAILS[idx % len(EMAILS)],
        "school": SCHOOLS[idx % len(SCHOOLS)],
        "degree": DEGREES[idx % len(DEGREES)],
        "edu_date": EDU_DATES[idx % len(EDU_DATES)],
        "works": works,
        "certs": sorted(set(certs)),
    }


def ground_truth(p: dict) -> dict:
    return {
        "name": p["name"], "phone": p["phone"], "email": p["email"],
        "schools": [p["school"]], "degrees": [p["degree"]],
        "companies": [w["company"] for w in p["works"]],
        "certificates": p["certs"],
        "date_ranges": sorted(set([p["edu_date"]] + [w["date"] for w in p["works"]])),
        "date_ranges_edu": p["edu_date"],
        "_works": [{"company": w["company"], "date": w["date"], "desc": w["desc"]} for w in p["works"]],
    }


# ---------------- PDF 生成 ----------------

def _register_cn_font() -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))


def gen_single_col(path: Path, p: dict) -> None:
    from reportlab.pdfgen import canvas

    _register_cn_font()

    c = canvas.Canvas(str(path), pagesize=(595, 842))
    y = 800
    c.setFont("STSong-Light", 16)
    c.drawString(50, y, p["name"])
    y -= 22
    c.setFont("STSong-Light", 10)
    c.drawString(50, y, f"电话: {p['phone']}   邮箱: {p['email']}")
    y -= 28
    for section, items in [
        ("教育经历", [f"{p['school']} | {p['degree']} | {p['edu_date']}"]),
        ("实习经历", [f"{w['company']} | {w['title']} | {w['date']}\n{w['desc']}" for w in p["works"]]),
        ("技能证书", ["、".join(p["certs"])]),
    ]:
        c.setFont("STSong-Light", 12)
        c.drawString(50, y, section)
        y -= 18
        c.setFont("STSong-Light", 10)
        for it in items:
            for ln in it.split("\n"):
                c.drawString(50, y, ln)
                y -= 15
        y -= 12
    c.save()


def gen_two_col(path: Path, p: dict) -> None:
    """双栏：左窄栏（联系方式/教育/证书/时间轴），右宽栏（经历：公司+日期+职责）。

    绘制顺序按**视觉行序交错**（左一行右一行）——这是设计模板 PDF 内容流的
    真实形态（也是问题报告 RPI-001 的机制：内容流顺序 ≠ 阅读顺序）。
    线性抽取因此把左右栏逐行混排，公司行与职责行的邻近关系被打乱。
    """
    from reportlab.pdfgen import canvas

    _register_cn_font()
    c = canvas.Canvas(str(path), pagesize=(595, 842))
    c.setFont("STSong-Light", 16)
    c.drawString(50, 800, p["name"])  # 页眉通栏姓名

    # 先把两栏内容各自排好行
    left_lines: list[tuple[float, int, str]] = []
    y = 760
    left_lines.append((y, 9, f"电话: {p['phone']}")); y -= 14
    left_lines.append((y, 9, f"邮箱: {p['email']}")); y -= 24
    left_lines.append((y, 11, "教育")); y -= 16
    left_lines.append((y, 9, p["school"])); y -= 13
    left_lines.append((y, 9, f"{p['degree']} {p['edu_date']}")); y -= 24
    left_lines.append((y, 11, "证书")); y -= 16
    left_lines.append((y, 9, "、".join(p["certs"]))); y -= 24
    left_lines.append((y, 11, "时间轴")); y -= 16
    for w in p["works"]:
        left_lines.append((y, 9, f"{w['date']} {w['company']}")); y -= 14

    right_lines: list[tuple[float, int, str]] = []
    ry = 760
    right_lines.append((ry, 12, "实习经历")); ry -= 18
    for w in p["works"]:
        right_lines.append((ry, 10, f"{w['company']} {w['title']}")); ry -= 15
        right_lines.append((ry, 9, w["date"])); ry -= 13
        right_lines.append((ry, 9, w["desc"])); ry -= 22

    # 按视觉行序交错绘制（复刻模板内容流≠阅读顺序的形态）
    merged = sorted(left_lines + right_lines, key=lambda t: -t[0])
    for ly, size, text in merged:
        x = 50 if (ly, size, text) in left_lines else 330
        c.setFont("STSong-Light", size)
        c.drawString(x, ly, text)
    c.save()


def gen_scanned(path: Path, p: dict) -> None:
    """图片型 PDF：位图整页（无文本层）。"""
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    _register_cn_font()

    img = Image.new("RGB", (1190, 1684), "white")
    d = ImageDraw.Draw(img)
    try:
        font_big = ImageFont.truetype("/System/Library/Fonts/PingFang.ttc", 36)
        font = ImageFont.truetype("/System/Library/Fonts/PingFang.ttc", 22)
    except OSError:
        font_big = font = ImageFont.load_default()
    y = 60
    d.text((50, y), p["name"], fill="black", font=font_big); y += 56
    d.text((50, y), f"电话 {p['phone']}  邮箱 {p['email']}", fill="black", font=font); y += 44
    d.text((50, y), "教育经历", fill="black", font=font_big); y += 44
    d.text((50, y), f"{p['school']} {p['degree']} {p['edu_date']}", fill="black", font=font); y += 44
    d.text((50, y), "实习经历", fill="black", font=font_big); y += 44
    for w in p["works"]:
        d.text((50, y), f"{w['company']} {w['title']} {w['date']}", fill="black", font=font); y += 36
        d.text((50, y), w["desc"], fill="black", font=font); y += 36
    d.text((50, y), "技能证书 " + "、".join(p["certs"]), fill="black", font=font)

    c = canvas.Canvas(str(path), pagesize=(595, 842))
    c.drawImage(ImageReader(img), 0, 0, width=595, height=842)
    c.save()




def gen_docx(path: Path, p: dict) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()
    doc.add_heading(p["name"], level=1)
    doc.add_paragraph(f"电话: {p['phone']}  邮箱: {p['email']}")
    doc.add_heading("教育经历", level=2)
    t = doc.add_table(rows=2, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.cell(0, 0).merge(t.cell(0, 2)).text = p["school"]  # 合并单元格：学校横跨
    t.cell(1, 0).text = p["degree"]
    t.cell(1, 1).text = p["edu_date"]
    t.cell(1, 2).text = "电子商务类"
    doc.add_heading("实习经历", level=2)
    for w in p["works"]:
        doc.add_paragraph(f"{w['company']} | {w['title']} | {w['date']}", style="List Bullet")
        doc.add_paragraph(w["desc"])
    doc.add_heading("技能证书", level=2)
    doc.add_paragraph("、".join(p["certs"]))
    doc.save(str(path))


def main() -> None:
    outdir = Path(sys.argv[sys.argv.index("--outdir") + 1]) if "--outdir" in sys.argv else Path("evals")
    corpus = outdir / "corpus"
    gt_dir = outdir / "ground_truth"
    for d in (corpus, gt_dir):
        d.mkdir(parents=True, exist_ok=True)

    rng = random.Random(20261013)
    profiles = [make_profile(i, rng) for i in range(10)]

    for i, p in enumerate(profiles, 1):
        gen_single_col(corpus / f"single-{i:02d}.pdf", p)
        (gt_dir / f"single-{i:02d}.json").write_text(
            json.dumps(ground_truth(p), ensure_ascii=False, indent=1), encoding="utf-8")
    for i, p in enumerate(profiles, 1):
        gen_two_col(corpus / f"two-col-{i:02d}.pdf", p)
        (gt_dir / f"two-col-{i:02d}.json").write_text(
            json.dumps(ground_truth(p), ensure_ascii=False, indent=1), encoding="utf-8")
    for i, p in enumerate(profiles, 1):
        gen_scanned(corpus / f"scanned-{i:02d}.pdf", p)
        (gt_dir / f"scanned-{i:02d}.json").write_text(
            json.dumps(ground_truth(p), ensure_ascii=False, indent=1), encoding="utf-8")
    for i, p in enumerate(profiles[:5], 1):
        gen_docx(corpus / f"docx-{i:02d}.docx", p)
        (gt_dir / f"docx-{i:02d}.json").write_text(
            json.dumps(ground_truth(p), ensure_ascii=False, indent=1), encoding="utf-8")

    total = len(list(corpus.iterdir()))
    print(f"语料 {total} 份 → {corpus}")
    print(f"ground truth {len(list(gt_dir.iterdir()))} 份 → {gt_dir}")


if __name__ == "__main__":
    main()
